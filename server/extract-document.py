import io
import sys
from html.parser import HTMLParser


class TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)


def main():
    extension = sys.argv[1].lower()
    raw = sys.stdin.buffer.read()

    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif extension == ".docx":
        from docx import Document
        document = Document(io.BytesIO(raw))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    elif extension == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(raw))
        text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    elif extension == ".html":
        parser = TextExtractor()
        parser.feed(raw.decode("utf-8", errors="replace"))
        text = "\n".join(parser.parts)
    else:
        text = raw.decode("utf-8-sig", errors="replace")

    sys.stdout.write(text)


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        sys.stderr.write(f"Could not read document: {error}")
        sys.exit(1)
